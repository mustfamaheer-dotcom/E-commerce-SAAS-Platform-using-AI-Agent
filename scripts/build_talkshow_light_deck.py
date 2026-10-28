# scripts/build_talkshow_light_deck.py
"""
Build the Graduation Project Presentation for TalkShow Platform.
Produces an ultra-modern tech Light-Themed PowerPoint deck (16:9 widescreen)
with heavy visual emphasis: embedding high-resolution charts, architecture diagrams,
and data science benchmarks while keeping text punchy and concise.
"""
import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Color Palette Constants (Ultra-Modern Tech Light Theme)
# -----------------------------------------------------------------------------
BG_SLIDE = RGBColor(248, 250, 252)      # #F8FAFC Slate-50 (Crisp clean background)
BG_CARD = RGBColor(255, 255, 255)       # #FFFFFF Pure White card
BG_CARD_SUBTLE = RGBColor(241, 245, 249)# #F1F5F9 Slate-100 container
BG_CARD_ACCENT = RGBColor(240, 249, 255)# #F0F9FF Sky-50 container

BORDER_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0 Subtle Slate border
BORDER_MUTED = RGBColor(203, 213, 225)  # #CBD5E1 Slate-300 border
BORDER_FOCUS = RGBColor(148, 163, 184)  # #94A3B8

# High-contrast readable typography for Light Theme
TEXT_DARK = RGBColor(15, 23, 42)        # #0F172A Slate-900 (Main Title)
TEXT_HEADING = RGBColor(30, 41, 59)     # #1E293B Slate-800
TEXT_BODY = RGBColor(51, 65, 85)        # #334155 Slate-700
TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B Slate-500
TEXT_LIGHT = RGBColor(148, 163, 184)    # #94A3B8 Slate-400

# Vibrant Modern Tech Accents
PRIMARY_BLUE = RGBColor(2, 132, 199)    # #0284C7 Sky-600 Primary
PRIMARY_DARK = RGBColor(3, 105, 161)    # #0369A1 Sky-700
PRIMARY_BG = RGBColor(224, 242, 254)    # #E0F2FE Sky-100 Pill fill

EMERALD_GREEN = RGBColor(5, 150, 105)   # #059669 Emerald-600
EMERALD_DARK = RGBColor(6, 95, 70)      # #065F46 Emerald-800
EMERALD_BG = RGBColor(209, 250, 229)    # #D1FAE5 Emerald-100

AI_VIOLET = RGBColor(124, 58, 237)      # #7C3AED Violet-600
AI_VIOLET_DARK = RGBColor(91, 33, 182)  # #5B21B6 Violet-800
AI_VIOLET_BG = RGBColor(237, 233, 254)  # #EDE9FE Violet-100

AMBER_ORANGE = RGBColor(217, 119, 6)    # #D97706 Amber-600
AMBER_DARK = RGBColor(146, 64, 14)      # #92400E Amber-800
AMBER_BG = RGBColor(254, 243, 199)      # #FEF3C7 Amber-100

ROSE_RED = RGBColor(225, 29, 72)        # #E11D48 Rose-600
ROSE_DARK = RGBColor(159, 18, 57)       # #9F1239 Rose-800
ROSE_BG = RGBColor(255, 228, 230)       # #FFE4E6 Rose-100

LOGO_TRANSPARENT = r"D:\HITU\2\AI\Project-AI\logo\TransparentLogo.png"
LOGO_SOLID = r"D:\HITU\2\AI\Project-AI\logo\logo.png"

ASSETS_DIR = r"D:\HITU\2\AI\Project-AI\Docs\Presentation\assets"
OUTPUT_DIR = r"D:\HITU\2\AI\Project-AI\Docs"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "TalkShow_Graduation_Presentation.pptx")
OUTPUT_PRES_DIR = r"D:\HITU\2\AI\Project-AI\Docs\Presentation"
OUTPUT_PRES_FILE = os.path.join(OUTPUT_PRES_DIR, "TalkShow_Graduation_Presentation.pptx")

# -----------------------------------------------------------------------------
# Helper Functions
# -----------------------------------------------------------------------------
def set_slide_background(slide, color=BG_SLIDE):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title, category_badge="TALKSHOW ARCHITECTURE", subtitle=None, badge_color=PRIMARY_DARK, badge_bg=PRIMARY_BG, badge_border=PRIMARY_BLUE):
    logo_file = LOGO_TRANSPARENT if os.path.exists(LOGO_TRANSPARENT) else LOGO_SOLID
    if os.path.exists(logo_file):
        try:
            slide.shapes.add_picture(logo_file, Inches(11.0), Inches(0.35), height=Inches(0.68))
        except Exception:
            pass

    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.38), Inches(3.5), Inches(0.32))
    pill.fill.solid()
    pill.fill.fore_color.rgb = badge_bg
    pill.line.color.rgb = badge_border
    pill.line.width = Pt(1)
    tf_pill = pill.text_frame
    tf_pill.word_wrap = True
    tf_pill.margin_left = tf_pill.margin_top = tf_pill.margin_bottom = tf_pill.margin_right = 0
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = f"●  {category_badge.upper()}"
    p_pill.font.name = "Segoe UI"
    p_pill.font.size = Pt(9)
    p_pill.font.bold = True
    p_pill.font.color.rgb = badge_color
    p_pill.alignment = PP_ALIGN.CENTER

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(10.0), Inches(0.72))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Segoe UI"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED

def add_footer(slide, current_slide, total_slides=16):
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.98), Inches(11.733), Inches(0.015))
    div.fill.solid()
    div.fill.fore_color.rgb = BORDER_LIGHT
    div.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = f"TalkShow Agentic E-Commerce Platform | Graduation Project 2026 — Data Science Department               Slide {current_slide} of {total_slides}"
    p.font.name = "Segoe UI"
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_MUTED

def add_modern_card(slide, left, top, width, height, bg_color=BG_CARD, border_color=BORDER_LIGHT, top_stripe_color=None, stripe_height=0.06):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()

    if top_stripe_color:
        stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(stripe_height))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = top_stripe_color
        stripe.line.fill.background()

    return card

def add_picture_safe(slide, img_path, left, top, width=None, height=None):
    if os.path.exists(img_path):
        try:
            if width and height:
                return slide.shapes.add_picture(img_path, left, top, width=width, height=height)
            elif width:
                return slide.shapes.add_picture(img_path, left, top, width=width)
            elif height:
                return slide.shapes.add_picture(img_path, left, top, height=height)
            else:
                return slide.shapes.add_picture(img_path, left, top)
        except Exception as e:
            print(f"Warning: Failed to load picture {img_path}: {e}")
    return None

def set_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

# -----------------------------------------------------------------------------
# Main Slide Deck Generator
# -----------------------------------------------------------------------------
def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Cover - Modern Light Tech)
    # =========================================================================
    s1 = prs.slides.add_slide(layout)
    set_slide_background(s1, BG_SLIDE)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_BLUE
    bar.line.fill.background()

    add_modern_card(s1, Inches(0.8), Inches(1.1), Inches(7.6), Inches(5.6), BG_CARD, BORDER_LIGHT, top_stripe_color=PRIMARY_BLUE)

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.35), Inches(6.8), Inches(5.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "● FACULTY OF COMPUTERS & ARTIFICIAL INTELLIGENCE | GRADUATION PROJECT 2026"
    p_badge.font.name = "Segoe UI"
    p_badge.font.size = Pt(10)
    p_badge.font.bold = True
    p_badge.font.color.rgb = PRIMARY_DARK

    p_title = tf.add_paragraph()
    p_title.text = "TalkShow"
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_DARK
    p_title.space_before = Pt(8)

    p_sub = tf.add_paragraph()
    p_sub.text = "Autonomous Agentic E-Commerce Platform & Conversational ERP"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(17)
    p_sub.font.bold = True
    p_sub.font.color.rgb = PRIMARY_BLUE
    p_sub.space_before = Pt(2)

    p_desc = tf.add_paragraph()
    p_desc.text = "Transforming Egyptian & MENA retail operations through LangGraph multi-agent orchestration, AST-sandboxed Text-to-SQL, quantile demand forecasting, and Human-in-the-Loop safety."
    p_desc.font.name = "Segoe UI"
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = TEXT_BODY
    p_desc.space_before = Pt(10)

    p_div = tf.add_paragraph()
    p_div.text = "―" * 44
    p_div.font.size = Pt(9)
    p_div.font.color.rgb = BORDER_LIGHT
    p_div.space_before = Pt(8)

    p_team_h = tf.add_paragraph()
    p_team_h.text = "Engineering & Data Science Team:"
    p_team_h.font.name = "Segoe UI"
    p_team_h.font.size = Pt(11)
    p_team_h.font.bold = True
    p_team_h.font.color.rgb = TEXT_HEADING
    p_team_h.space_before = Pt(6)

    p_team = tf.add_paragraph()
    p_team.text = "مصطفى ماهر  •  عادل أحمد  •  عبد الله سامح  •  ريتاج حكيم  •  رناد حكيم  •  رُؤى أشرف"
    p_team.font.name = "Segoe UI"
    p_team.font.size = Pt(10.5)
    p_team.font.bold = True
    p_team.font.color.rgb = PRIMARY_DARK
    p_team.space_before = Pt(2)

    p_sup = tf.add_paragraph()
    p_sup.text = "Supervised by: Prof. Dr. Ahmed El-Sayed & Academic Committee"
    p_sup.font.name = "Segoe UI"
    p_sup.font.size = Pt(10)
    p_sup.font.color.rgb = TEXT_MUTED
    p_sup.space_before = Pt(6)

    logo_file = LOGO_TRANSPARENT if os.path.exists(LOGO_TRANSPARENT) else LOGO_SOLID
    if os.path.exists(logo_file):
        try:
            s1.shapes.add_picture(logo_file, Inches(8.8), Inches(1.8), width=Inches(3.8))
        except Exception:
            pass

    set_speaker_notes(s1, """ملاحظات المتحدث (Slide 1 - Cover):
بسم الله الرحمن الرحيم، أساتذتنا المشرفين وأعضاء لجنة التحكيم الموقرة، يسعدنا ويشرفنا تقديم مشروع تخرجنا: منصة توك شو (TalkShow).
المنصة بتمثل نقلة نوعية في التجارة الإلكترونية وأنظمة إدارة الموارد (ERP)، حيث بنحول العمليات اليومية المعقدة لتجربة محادثة ذكية متصلة بقواعد البيانات بنظام متعدد الوكلاء (Multi-Agent System).
فريق العمل مكون من 6 باحثين ومهندسين وزعنا التخصصات بدقة بين معمارية النظم، ووكلاء LangGraph، وتدريب نماذج تعلم الآلة، وحماية قواعد البيانات.""")

    # =========================================================================
    # SLIDE 2: EXECUTIVE SUMMARY & THE SYSTEM VISION
    # =========================================================================
    s2 = prs.slides.add_slide(layout)
    set_slide_background(s2, BG_SLIDE)
    add_header(s2, "Executive Summary: Paradigm Shift in E-Commerce & ERP", "PLATFORM OVERVIEW & THESIS", "Why traditional enterprise software is obsolete and how TalkShow transforms business operations")
    add_footer(s2, 2)

    pillars = [
        ("Autonomous Conversational Execution", PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG, [
            ("Replaces 50+ Admin Screens", "Merchants execute operations via natural Egyptian dialogue."),
            ("Human-in-the-Loop (HITL)", "Mutations run dry-run simulations before cryptographic commit."),
            ("Generative UI (GenUI)", "Emits dynamic React cards (<BulkPricePreviewCard />) instead of text.")
        ]),
        ("Grounded Dialectal Intelligence", EMERALD_GREEN, EMERALD_DARK, EMERALD_BG, [
            ("Egyptian Colloquial Arabic", "Native parsing for retail slang ('السيلز', 'البرودكتس') and typos."),
            ("Zero Price Hallucination", "Strict verification against live PostgreSQL database records."),
            ("Local Ecosystem Connectors", "Built-in Paymob payment checkout and Bosta courier logistics.")
        ]),
        ("Predictive MLOps & Conversational BI", AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG, [
            ("Safe Text-to-SQL", "Translates questions to AST-sandboxed SQL across 71+ tables."),
            ("Quantile Demand Forecasting", "LightGBM models deliver 30-day forecast bands (95% CI)."),
            ("Dynamic Stockout Prevention", "Algorithmic safety stock and reorder points (ROP) calculation.")
        ])
    ]

    left_positions = [Inches(0.8), Inches(4.8), Inches(8.8)]
    for i, (title, color, dark_color, bg_pill, bullets) in enumerate(pillars):
        add_modern_card(s2, left_positions[i], Inches(1.55), Inches(3.73), Inches(3.55), BG_CARD, BORDER_LIGHT, top_stripe_color=color)

        card_pill = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_positions[i] + Inches(0.2), Inches(1.75), Inches(3.33), Inches(0.35))
        card_pill.fill.solid()
        card_pill.fill.fore_color.rgb = bg_pill
        card_pill.line.color.rgb = color
        card_pill.line.width = Pt(1)
        tf_cp = card_pill.text_frame
        p_cp = tf_cp.paragraphs[0]
        p_cp.text = title.upper()
        p_cp.font.name = "Segoe UI"
        p_cp.font.size = Pt(8.5)
        p_cp.font.bold = True
        p_cp.font.color.rgb = dark_color
        p_cp.alignment = PP_ALIGN.CENTER

        tb = s2.shapes.add_textbox(left_positions[i] + Inches(0.2), Inches(2.2), Inches(3.33), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True
        for b_title, b_desc in bullets:
            p_bt = tf.add_paragraph()
            p_bt.text = f"✔ {b_title}"
            p_bt.font.name = "Segoe UI"
            p_bt.font.size = Pt(10.5)
            p_bt.font.bold = True
            p_bt.font.color.rgb = TEXT_HEADING
            p_bt.space_before = Pt(6)

            p_bd = tf.add_paragraph()
            p_bd.text = b_desc
            p_bd.font.name = "Segoe UI"
            p_bd.font.size = Pt(9)
            p_bd.font.color.rgb = TEXT_BODY
            p_bd.space_before = Pt(1)

    # Bottom Stat Badges
    stat_boxes = [
        ("70%", "Token Optimization", "Dynamic vector schema pruner", PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("100%", "Security Block Rate", "sqlglot AST sandbox defense", ROSE_RED, ROSE_DARK, ROSE_BG),
        ("14.2%", "WAPE Forecast Error", "Beats Prophet (28.6%) & ARIMA", EMERALD_GREEN, EMERALD_DARK, EMERALD_BG),
        ("2.18 GB", "Edge VRAM Footprint", "Runs air-gapped on RTX 3050", AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG)
    ]
    for i, (metric, label, sub, color, dark_color, bg_pill) in enumerate(stat_boxes):
        x = Inches(0.8 + i * 2.98)
        add_modern_card(s2, x, Inches(5.25), Inches(2.8), Inches(1.5), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s2.shapes.add_textbox(x + Inches(0.15), Inches(5.35), Inches(2.5), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p_m = tf.paragraphs[0]
        p_m.text = metric
        p_m.font.name = "Segoe UI"
        p_m.font.size = Pt(24)
        p_m.font.bold = True
        p_m.font.color.rgb = dark_color
        p_m.alignment = PP_ALIGN.CENTER

        p_l = tf.add_paragraph()
        p_l.text = label
        p_l.font.name = "Segoe UI"
        p_l.font.size = Pt(9.5)
        p_l.font.bold = True
        p_l.font.color.rgb = TEXT_HEADING
        p_l.alignment = PP_ALIGN.CENTER
        p_l.space_before = Pt(1)

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(8)
        p_s.font.color.rgb = TEXT_MUTED
        p_s.alignment = PP_ALIGN.CENTER

    set_speaker_notes(s2, """ملاحظات المتحدث (Slide 2 - Executive Summary):
الفكرة الجوهرية لمنصة توك شو هي القضاء على الفجوة الكبيرة بين تعقيد أنظمة الـ ERP التقليدية والتجار الحقيقيين.
التاجر يطلب أمر باللغة الطبيعية والوكيل الذكي ينفذه بأمان تام بعد مراجعة هامش الربح وعرض بطاقة تفاعلية GenUI للموافقة.
بنوفر 3 ركائز: تنفيذ آمن مع موافقة بشرية، فهم عميق للهجة المصرية وقوانين التجارة المحلية، ونماذج تعلم آلي استباقية تتنبأ بنفاد المخزون وتفصل العملاء لشرائح استهلاكية.""")

    # =========================================================================
    # SLIDE 3: SECTION 1 — THE PROBLEM: DEEP INDUSTRY REALITIES
    # =========================================================================
    s3 = prs.slides.add_slide(layout)
    set_slide_background(s3, BG_SLIDE)
    add_header(s3, "Section 1: The Core Industry Problem & Market Realities", "1. PROBLEM IDENTIFICATION", "Operational fragmentation, data silos, and the retail crisis in Egyptian & MENA SMEs", badge_color=ROSE_DARK, badge_bg=ROSE_BG, badge_border=ROSE_RED)
    add_footer(s3, 3)

    # Left: 2 Concise Problem Cards
    add_modern_card(s3, Inches(0.8), Inches(1.55), Inches(4.5), Inches(2.55), BG_CARD, BORDER_LIGHT, top_stripe_color=ROSE_RED)
    tb_p1 = s3.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.1), Inches(2.2))
    tf_p1 = tb_p1.text_frame
    tf_p1.word_wrap = True
    p1 = tf_p1.paragraphs[0]
    p1.text = "1. Operational Fragmentation"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = ROSE_DARK
    for b in [
        "Merchants juggle 5+ tools: Excel sheets, POS, and WhatsApp.",
        "42% of daily workday wasted on manual data entry.",
        "Frequent inventory overselling due to manual lag."
    ]:
        p = tf_p1.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    add_modern_card(s3, Inches(0.8), Inches(4.2), Inches(4.5), Inches(2.55), BG_CARD, BORDER_LIGHT, top_stripe_color=AMBER_ORANGE)
    tb_p2 = s3.shapes.add_textbox(Inches(1.0), Inches(4.35), Inches(4.1), Inches(2.2))
    tf_p2 = tb_p2.text_frame
    tf_p2.word_wrap = True
    p2 = tf_p2.paragraphs[0]
    p2.text = "2. Relational BI Blindspot & Stockout"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = AMBER_DARK
    for b in [
        "Data trapped in 71+ tables inaccessible to business merchants.",
        "23.4% annual revenue lost to stockouts in Egyptian SMEs.",
        "Zero proactive warning before inventory depletes below safety stock."
    ]:
        p = tf_p2.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(4)

    # Right: Chart Picture
    add_modern_card(s3, Inches(5.5), Inches(1.55), Inches(7.033), Inches(5.2), BG_CARD, BORDER_LIGHT)
    chart_problem = os.path.join(ASSETS_DIR, "chart_problem_time_loss.png")
    add_picture_safe(s3, chart_problem, Inches(5.6), Inches(1.7), width=Inches(6.833))

    set_speaker_notes(s3, """ملاحظات المتحدث (Slide 3 - Problem):
لما حللنا واقع التجارة الإلكترونية والتجزئة في مصر والمنطقة العربية، اكتشفنا أزمتين حاسمتين توضحهما الرسوم البيانية:
الرسم الأول: التاجر بيضيع 42% من يومه في إدخال ومطابقة البيانات يدوياً بين الإكسيل والواتساب، وفقط 12% للتطوير الاستراتيجي.
الرسم الثاني: الشركات المصرية بتخسر 23.4% من مبيعاتها السنوية بسبب نفاد المخزون مقارنة بالمعدل العالمي (4.5%)، لأن البيانات محبوسة في جداول الـ ERP وما عندهمش أدوات استباقية.""")

    # =========================================================================
    # SLIDE 4: SECTION 1 — PROBLEM DEEP DIVE: RISKS OF AUTONOMOUS AI
    # =========================================================================
    s4 = prs.slides.add_slide(layout)
    set_slide_background(s4, BG_SLIDE)
    add_header(s4, "Problem Deep Dive: Why Generic LLMs Fail in Enterprise Operations", "1. DEEP PROBLEM ANALYSIS", "The fatal flaws of naive LLM wrappers in mission-critical database & e-commerce systems", badge_color=ROSE_DARK, badge_bg=ROSE_BG, badge_border=ROSE_RED)
    add_footer(s4, 4)

    # Top Warning Badges
    top_risks = [
        ("Hallucination Liability", "Prices must comply with Egyptian Consumer Protection Law 181/2018.", ROSE_RED, ROSE_DARK, ROSE_BG),
        ("Unbounded SQL Mutations", "Direct write access (`UPDATE`/`DELETE`) can corrupt catalogs.", AMBER_ORANGE, AMBER_DARK, AMBER_BG),
        ("Multi-Tenant Data Spillage", "Flawed prompts leak competitor data in shared database tables.", AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG)
    ]
    for i, (title, sub, color, dark_color, bg_pill) in enumerate(top_risks):
        x = Inches(0.8 + i * 3.98)
        add_modern_card(s4, x, Inches(1.55), Inches(3.78), Inches(0.9), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s4.shapes.add_textbox(x + Inches(0.15), Inches(1.65), Inches(3.48), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = f"⚠ {title}"
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(8.5)
        p_s.font.color.rgb = TEXT_BODY
        p_s.space_before = Pt(2)

    # Bottom Visual Comparison Diagram
    add_modern_card(s4, Inches(0.8), Inches(2.6), Inches(11.733), Inches(4.15), BG_CARD, BORDER_LIGHT)
    diag_comp = os.path.join(ASSETS_DIR, "diagram_generic_vs_talkshow.png")
    add_picture_safe(s4, diag_comp, Inches(1.0), Inches(2.7), width=Inches(11.333))

    set_speaker_notes(s4, """ملاحظات المتحدث (Slide 4 - Deep Problem Analysis):
هنا في قسم علوم البيانات كان لازم نسأل سؤال علمي حاسم: ليه ما نربطش ChatGPT بقاعدة البيانات وخلاص؟
الرسم التوضيحي بيقارن بين الـ Naive LLM Wrapper ومعمارية توك شو:
1- الهلوسة السعرية: لو الموديل وعد العميل بخصم غير موجود، التاجر بيتحمل مسؤولية قانونية بموجب قانون حماية المستهلك المصري.
2- عمليات الكتابة المفتوحة: منح الذكاء الاصطناعي صلاحية تعديل مباشرة خطر جداً.
3- الحل في توك شو: جدار حماية AST، محاكاة جافة Dry-Run، وعزل المستأجرين على مستوى نواة PostgreSQL RLS.""")

    # =========================================================================
    # SLIDE 5: SECTION 2 — KEY ENGINEERING & DATA SCIENCE CHALLENGES
    # =========================================================================
    s5 = prs.slides.add_slide(layout)
    set_slide_background(s5, BG_SLIDE)
    add_header(s5, "Section 2: Engineering & Data Science Challenges", "2. TECHNICAL CHALLENGES", "Overcoming relational complexity, non-deterministic reasoning, and multi-tenant security barriers", badge_color=AMBER_DARK, badge_bg=AMBER_BG, badge_border=AMBER_ORANGE)
    add_footer(s5, 5)

    challenges = [
        ("Dialectal NLP & Ambiguity",
         "Egyptian Arabic contains typos, local retail jargon ('السيلز', 'الأوردرات'), and missing diacritics.",
         "Solution: Multi-stage normalization and custom entity mapping pipeline before schema matching.",
         PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("Non-Deterministic LLM to Strict SQL",
         "LLMs can produce invalid joins, non-existent column names, or catastrophic data deletion statements.",
         "Solution: Abstract Syntax Tree (AST) validation via `sqlglot` strictly enforcing read-only execution.",
         ROSE_RED, ROSE_DARK, ROSE_BG),
        ("Multi-Tenant Kernel Isolation",
         "Application-level `WHERE tenant_id = x` filtering is prone to human omissions and prompt leaks.",
         "Solution: PostgreSQL 16 Row-Level Security (`app.current_tenant_id`) enforced at the database kernel.",
         AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG),
        ("Edge Hardware Constraints (4GB VRAM)",
         "SMEs require zero-cloud air-gapped privacy without paying recurring token API fees.",
         "Solution: 4-bit quantized Qwen2.5-Coder-3B (GGUF) running smoothly on consumer RTX 3050 GPUs (2.18 GB).",
         EMERALD_GREEN, EMERALD_DARK, EMERALD_BG)
    ]

    coords_ch = [
        (Inches(0.8), Inches(1.55)),
        (Inches(6.8), Inches(1.55)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.8), Inches(4.2))
    ]

    for (title, prob, sol, color, dark_color, bg_pill), (x, y) in zip(challenges, coords_ch):
        add_modern_card(s5, x, y, Inches(5.733), Inches(2.55), BG_CARD, BORDER_LIGHT, top_stripe_color=color)

        pill = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y + Inches(0.2), Inches(5.333), Inches(0.35))
        pill.fill.solid()
        pill.fill.fore_color.rgb = bg_pill
        pill.line.color.rgb = color
        pill.line.width = Pt(1)
        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = f"CHALLENGE: {title.upper()}"
        p_p.font.name = "Segoe UI"
        p_p.font.size = Pt(8.5)
        p_p.font.bold = True
        p_p.font.color.rgb = dark_color
        p_p.alignment = PP_ALIGN.CENTER

        tb = s5.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), Inches(5.333), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p_pr = tf.paragraphs[0]
        p_pr.text = f"Issue: {prob}"
        p_pr.font.name = "Segoe UI"
        p_pr.font.size = Pt(9.5)
        p_pr.font.color.rgb = TEXT_BODY

        p_sl = tf.add_paragraph()
        p_sl.text = f"✔ {sol}"
        p_sl.font.name = "Segoe UI"
        p_sl.font.size = Pt(9.5)
        p_sl.font.bold = True
        p_sl.font.color.rgb = TEXT_HEADING
        p_sl.space_before = Pt(6)

    set_speaker_notes(s5, """ملاحظات المتحدث (Slide 5 - Technical Challenges):
واجهنا 4 تحديات هندسية وعلمية رئيسية:
1- اللهجة العامية والفرانكو: حليناها بخوارزمية تنقية وتحويل لغوي تترجم الكلمات الدارجة لمفردات المخطط.
2- عدم حتمية الـ LLMs: حليناها بجدار AST يرفض أي استعلام خاطئ أو غير مصرح به قبل وصوله للداتابيز.
3- الأمان المتعدد للمستأجرين: وظفنا PostgreSQL Row-Level Security في نواة قاعدة البيانات نفسها.
4- قيود العتاد: قمنا بعمل كمية وضغط لنموذج مفتوح المصدر ليعمل محلياً باستهلاك 2.2 جيجا رام فقط.""")

    # =========================================================================
    # SLIDE 6: SECTION 3 — THE TALKSHOW SOLUTION: ARCHITECTURAL BLUEPRINT
    # =========================================================================
    s6 = prs.slides.add_slide(layout)
    set_slide_background(s6, BG_SLIDE)
    add_header(s6, "Section 3: The TalkShow Solution — Architectural Blueprint", "3. ARCHITECTURAL SOLUTION", "End-to-end data flow: from natural language input to verified database execution and Generative UI")
    add_footer(s6, 6)

    # Top Architectural Guarantees
    arch_badges = [
        ("Layered Architecture", "5-Tier production monorepo", PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("PromptGuard & PII", "Redacts Egyptian National IDs", EMERALD_GREEN, EMERALD_DARK, EMERALD_BG),
        ("HITL State Machine", "Deterministic LangGraph safety", AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG),
        ("Kernel RLS Isolation", "PostgreSQL 16 multi-tenant data", ROSE_RED, ROSE_DARK, ROSE_BG)
    ]
    for i, (title, sub, color, dark_color, bg_pill) in enumerate(arch_badges):
        x = Inches(0.8 + i * 2.98)
        add_modern_card(s6, x, Inches(1.55), Inches(2.8), Inches(0.85), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s6.shapes.add_textbox(x + Inches(0.15), Inches(1.65), Inches(2.5), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(8.5)
        p_s.font.color.rgb = TEXT_MUTED

    # Center Diagram Picture
    add_modern_card(s6, Inches(0.8), Inches(2.55), Inches(11.733), Inches(4.2), BG_CARD, BORDER_LIGHT)
    diag_arch = os.path.join(ASSETS_DIR, "diagram_architecture_flow.png")
    add_picture_safe(s6, diag_arch, Inches(1.0), Inches(2.65), width=Inches(11.333))

    set_speaker_notes(s6, """ملاحظات المتحدث (Slide 6 - Architectural Blueprint):
المعمارية الشاملة لمنصة توك شو مقسمة لـ 5 طبقات متكاملة توضحها الشاشة:
1- واجهة العرض: React 18 مع نظام Generative UI لبث بطاقات تفاعلية.
2- بوابة الـ API: جدار حماية PromptGuard وتتبع Latency لكل استدعاء.
3- محرك الوكلاء: LangGraph MasterCommandGraph يدير 10 مسارات وظيفية مع ذاكرة ثلاثية الأبعاد.
4- خادم الـ MLOps: نماذج LightGBM للتنبؤ بالطلب وتجزئة العملاء.
5- طبقة البيانات: PostgreSQL 16 RLS مع كاش Redis فائق السرعة.""")

    # =========================================================================
    # SLIDE 7: SECTION 3 — SOLUTION DEEP DIVE: LANGGRAPH & HITL GATE
    # =========================================================================
    s7 = prs.slides.add_slide(layout)
    set_slide_background(s7, BG_SLIDE)
    add_header(s7, "Solution Deep-Dive: Autonomous LangGraph Agent & HITL Safety", "3. AGENTIC ORCHESTRATION", "Deterministic state machines, multi-turn memory, and the Human-in-the-Loop approval gate", badge_color=AI_VIOLET_DARK, badge_bg=AI_VIOLET_BG, badge_border=AI_VIOLET)
    add_footer(s7, 7)

    # Top: 2 Summary Cards
    add_modern_card(s7, Inches(0.8), Inches(1.55), Inches(5.733), Inches(1.5), BG_CARD, BORDER_LIGHT, top_stripe_color=AI_VIOLET)
    tb_l = s7.shapes.add_textbox(Inches(1.0), Inches(1.68), Inches(5.333), Inches(1.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p_lt = tf_l.paragraphs[0]
    p_lt.text = "LangGraph MasterCommandGraph"
    p_lt.font.name = "Segoe UI"
    p_lt.font.size = Pt(12)
    p_lt.font.bold = True
    p_lt.font.color.rgb = AI_VIOLET_DARK
    for b in [
        "10 deterministic intent routes (replaces chaotic ReAct loops).",
        "3-Tier memory: short-term buffer, pgvector RAG, and merchant prefs."
    ]:
        p = tf_l.add_paragraph()
        p.text = f"✔ {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(2)

    add_modern_card(s7, Inches(6.8), Inches(1.55), Inches(5.733), Inches(1.5), BG_CARD, BORDER_LIGHT, top_stripe_color=PRIMARY_BLUE)
    tb_r = s7.shapes.add_textbox(Inches(7.0), Inches(1.68), Inches(5.333), Inches(1.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p_rt = tf_r.paragraphs[0]
    p_rt.text = "Human-in-the-Loop (HITL) Safety Gate"
    p_rt.font.name = "Segoe UI"
    p_rt.font.size = Pt(12)
    p_rt.font.bold = True
    p_rt.font.color.rgb = PRIMARY_DARK
    for b in [
        "Shadow Dry-Run simulates profit margin floor: Price > Cost.",
        "Generates cryptographic `actionId`; requires merchant click to commit."
    ]:
        p = tf_r.add_paragraph()
        p.text = f"▶ {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_BODY
        p.space_before = Pt(2)

    # Bottom Diagram Picture
    add_modern_card(s7, Inches(0.8), Inches(3.2), Inches(11.733), Inches(3.55), BG_CARD, BORDER_LIGHT)
    diag_hitl = os.path.join(ASSETS_DIR, "diagram_hitl_workflow.png")
    add_picture_safe(s7, diag_hitl, Inches(1.0), Inches(3.3), width=Inches(11.333))

    set_speaker_notes(s7, """ملاحظات المتحدث (Slide 7 - Agent Core & HITL):
الابتكار الأهم هنا هو بروتوكول الـ Human-in-the-Loop (HITL) الموضح في الرسم:
الوكيل الذكي ممنوع تماماً من الكتابة المباشرة في قاعدة البيانات؛ لما التاجر يطلب مثلاً: 'طبق خصم 15% على كل السماعات'، الوكيل بينفذ محاكاة جافة (Dry-Run)، ويحسب التكلفة وهامش الربح لكل منتج.
لو لقى إن الخصم هيبيع منتج بأقل من سعر التكلفة، بيحذر التاجر فوراً، ويصدر كارت GenUI تفاعلي فيه كل الفروق وزرار تأكيد. العملية مش بتتحفظ في الداتابيز إلا لما التاجر يضغط 'موافقة' بنفسه.""")

    # =========================================================================
    # SLIDE 8: SECTION 3 — SOLUTION DEEP DIVE: AST TEXT-TO-SQL
    # =========================================================================
    s8 = prs.slides.add_slide(layout)
    set_slide_background(s8, BG_SLIDE)
    add_header(s8, "Solution Deep-Dive: AST Sandboxed Text-to-SQL Engine", "3. CONVERSATIONAL BUSINESS INTELLIGENCE", "Translating Egyptian colloquial queries to AST-validated SQL across 71+ normalized tables")
    add_footer(s8, 8)

    # Top Highlights
    ast_badges = [
        ("Vector Schema Pruner", "Extracts only 3-5 relevant tables (70% token savings)", PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("sqlglot AST Sandbox", "Rejects DROP, DELETE, UPDATE, and catalog probes", ROSE_RED, ROSE_DARK, ROSE_BG),
        ("Row Clamping & Reflection", "LIMIT 500 enforced with automated reflection loop", EMERALD_GREEN, EMERALD_DARK, EMERALD_BG)
    ]
    for i, (title, sub, color, dark_color, bg_pill) in enumerate(ast_badges):
        x = Inches(0.8 + i * 3.98)
        add_modern_card(s8, x, Inches(1.55), Inches(3.78), Inches(0.85), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s8.shapes.add_textbox(x + Inches(0.15), Inches(1.65), Inches(3.48), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(8.5)
        p_s.font.color.rgb = TEXT_MUTED

    # Bottom AST Defense Diagram Picture
    add_modern_card(s8, Inches(0.8), Inches(2.55), Inches(11.733), Inches(4.2), BG_CARD, BORDER_LIGHT)
    diag_ast = os.path.join(ASSETS_DIR, "diagram_ast_defense.png")
    add_picture_safe(s8, diag_ast, Inches(1.0), Inches(2.65), width=Inches(11.333))

    set_speaker_notes(s8, """ملاحظات المتحدث (Slide 8 - AST Text-to-SQL):
محرك استعلام قواعد البيانات باللغة الطبيعية (Text-to-SQL) بيعتمد على جدار حماية خماسي الطبقات موضح بالرسم:
1- تطبيع اللهجة المصرية وفك الكلمات الدارجة.
2- استخلاص الجداول المرتبطة فقط بالمتجه Schema Pruning لتوفير 70% من التوكنز.
3- فحص شجرة الـ AST بمكتبة sqlglot ورفض أي أمر تعديل أو حذف.
4- حظر الفواصل المنقوطة لمنع Piggybacking.
5- فرض سقف 500 صف والتصحيح الذاتي في حالة الخطأ.""")

    # =========================================================================
    # SLIDE 9: SECTION 3 — SOLUTION DEEP DIVE: PREDICTIVE MLOPS & ML ENGINES
    # =========================================================================
    s9 = prs.slides.add_slide(layout)
    set_slide_background(s9, BG_SLIDE)
    add_header(s9, "Solution Deep-Dive: Predictive MLOps & Statistical ML", "3. MACHINE LEARNING & MLOPS", "Quantile gradient boosting, SHAP feature attributions, and algorithmic customer segmentation", badge_color=EMERALD_DARK, badge_bg=EMERALD_BG, badge_border=EMERALD_GREEN)
    add_footer(s9, 9)

    # Top: 3 Concise Capability Cards
    ml_caps = [
        ("Quantile Demand Forecast", "LightGBM regressor predicting 30-day sales with 95% CI bands. WAPE = 14.2% (vs Prophet 28.6%).", PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("SHAP Explainability", "TreeExplainer translates tree splits into plain retail drivers (e.g. Ramadan bump +34.2%, price cuts).", EMERALD_GREEN, EMERALD_DARK, EMERALD_BG),
        ("Log-RFM K-Means (K=5)", "Segments customer base into VIP, Loyal, New, At-Risk, and Churned for automated targeted win-backs.", AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG)
    ]
    for i, (title, desc, color, dark_color, bg_pill) in enumerate(ml_caps):
        x = Inches(0.8 + i * 3.98)
        add_modern_card(s9, x, Inches(1.55), Inches(3.78), Inches(1.15), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s9.shapes.add_textbox(x + Inches(0.15), Inches(1.65), Inches(3.48), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = TEXT_BODY
        p_d.space_before = Pt(2)

    # Bottom: 3 Charts Side-by-Side!
    chart_fc = os.path.join(ASSETS_DIR, "chart_forecast_cone.png")
    chart_sh = os.path.join(ASSETS_DIR, "chart_shap_importance.png")
    chart_rf = os.path.join(ASSETS_DIR, "chart_rfm_clusters.png")

    add_modern_card(s9, Inches(0.8), Inches(2.85), Inches(3.78), Inches(3.9), BG_CARD, BORDER_LIGHT)
    add_picture_safe(s9, chart_fc, Inches(0.85), Inches(2.95), width=Inches(3.68))

    add_modern_card(s9, Inches(4.78), Inches(2.85), Inches(3.78), Inches(3.9), BG_CARD, BORDER_LIGHT)
    add_picture_safe(s9, chart_sh, Inches(4.83), Inches(2.95), width=Inches(3.68))

    add_modern_card(s9, Inches(8.76), Inches(2.85), Inches(3.78), Inches(3.9), BG_CARD, BORDER_LIGHT)
    add_picture_safe(s9, chart_rf, Inches(8.81), Inches(2.95), width=Inches(3.68))

    set_speaker_notes(s9, """ملاحظات المتحدث (Slide 9 - Machine Learning Engines):
الجانب الإحصائي والتعلم الآلي في المشروع موضح بالأرقام والرسومات الفعلية:
1- الرسم الأول (يسار): مخروط التنبؤ الاحتمالي بالطلب لمدة 30 يوم بنسبة ثقة 95% مع خط إعادة الطلب ROP.
2- الرسم الأوسط: تحليل SHAP يوضح العوامل الحقيقية المؤثرة في الطلب مثل مواسم الأعياد وخصومات الأسعار.
3- الرسم الثالث (يمين): تجميع العملاء Log-RFM K-Means لـ 5 شرائح استهلاكية لإطلاق حملات استبقاء العملاء المعرضين للانسحاب.""")

    # =========================================================================
    # SLIDE 10: SECTION 4 — LIVE EXPERIENCE: GENERATIVE UI (GenUI)
    # =========================================================================
    s10 = prs.slides.add_slide(layout)
    set_slide_background(s10, BG_SLIDE)
    add_header(s10, "Section 4: Live Experience — Interactive Generative UI", "4. USER EXPERIENCE & INTERFACES", "Streaming dynamic React components instead of static text blocks for actionable operations")
    add_footer(s10, 10)

    # Left: Concise GenUI Card Highlights
    genui_specs = [
        ("<BulkPricePreviewCard />", "Displays before/after prices, profit margin delta, and 1-click commit button.", PRIMARY_BLUE, PRIMARY_DARK),
        ("<SalesForecastCard />", "Interactive Recharts 30-day trajectory with 95% confidence intervals.", EMERALD_GREEN, EMERALD_DARK),
        ("<InventoryAlertsList />", "Flags critical Days-to-Stockout (DUS) and drafts automated supplier Purchase Orders.", ROSE_RED, ROSE_DARK)
    ]
    for i, (tag, desc, color, dark_color) in enumerate(genui_specs):
        y = Inches(1.55 + i * 1.75)
        add_modern_card(s10, Inches(0.8), y, Inches(4.6), Inches(1.6), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s10.shapes.add_textbox(Inches(1.0), y + Inches(0.12), Inches(4.2), Inches(1.35))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = tag
        p_t.font.name = "Consolas"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = TEXT_BODY
        p_d.space_before = Pt(3)

    # Right: GenUI Pipeline Diagram
    add_modern_card(s10, Inches(5.6), Inches(1.55), Inches(6.933), Inches(5.2), BG_CARD, BORDER_LIGHT)
    diag_genui = os.path.join(ASSETS_DIR, "diagram_genui_concept.png")
    add_picture_safe(s10, diag_genui, Inches(5.75), Inches(1.9), width=Inches(6.633))

    set_speaker_notes(s10, """ملاحظات المتحدث (Slide 10 - Generative UI):
الرسم التوضيحي بيوضح فلسفة الـ Generative UI في توك شو:
بدل ما نرد على التاجر بفقرات نصية طويلة، الوكيل الذكي بيرسل Structured JSON بيترجم فوراً في واجهة React لبطاقات تفاعلية:
التاجر بيشوف جدول مقارنة الأسعار وهامش الربح، وبيعمل الموافقة أو التعديل من داخل نفس الكارت بنقرة واحدة.""")

    # =========================================================================
    # SLIDE 11: SECTION 4 — LIVE EXPERIENCE: GROUNDED EGYPTIAN DIALECT IN ACTION
    # =========================================================================
    s11 = prs.slides.add_slide(layout)
    set_slide_background(s11, BG_SLIDE)
    add_header(s11, "Live Experience: Real-World Dialectal NLP Scenarios", "4. LIVE SYSTEM DEMONSTRATION", "Executing complex multi-step operations through natural Egyptian colloquial conversation")
    add_footer(s11, 11)

    # Top: 2 Real-world Scenarios
    scenarios_top = [
        ("Scenario A: Bulk Discount with HITL",
         "\"طبق خصم 15% على السماعات اللي رصيدها اكتر من 50 قطعه\"",
         "• Normalizes 'سماعات' -> Category 'Headphones'\n• Shadow dry-run verifies margin floor: avg margin 38.5% -> 27.6%\n• Emits `<BulkPricePreviewCard />` with 1-click Approve button",
         PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("Scenario B: Predictive Stockout Alert",
         "\"قولي ايه البضاعة اللي ممكن تخلص مننا الأسبوع الجاي\"",
         "• Runs StockoutOptimizer across tenant inventory\n• Calculates Days Until Stockout: DUS = Q / EWMA_d <= 7 days\n• Emits `<InventoryAlertsList />` with instant Purchase Order draft",
         EMERALD_GREEN, EMERALD_DARK, EMERALD_BG)
    ]
    for i, (title, prompt, steps, color, dark_color, bg_pill) in enumerate(scenarios_top):
        x = Inches(0.8 + i * 5.95)
        add_modern_card(s11, x, Inches(1.55), Inches(5.75), Inches(1.75), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s11.shapes.add_textbox(x + Inches(0.18), Inches(1.68), Inches(5.4), Inches(1.55))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_pr = tf.add_paragraph()
        p_pr.text = f"Prompt: {prompt}"
        p_pr.font.name = "Segoe UI"
        p_pr.font.size = Pt(9.5)
        p_pr.font.bold = True
        p_pr.font.color.rgb = PRIMARY_DARK
        p_pr.space_before = Pt(2)

        p_st = tf.add_paragraph()
        p_st.text = steps
        p_st.font.name = "Segoe UI"
        p_st.font.size = Pt(8.5)
        p_st.font.color.rgb = TEXT_BODY
        p_st.space_before = Pt(2)

    # Bottom: Dialect Translation Pipeline Diagram
    add_modern_card(s11, Inches(0.8), Inches(3.45), Inches(11.733), Inches(3.3), BG_CARD, BORDER_LIGHT)
    diag_dialect = os.path.join(ASSETS_DIR, "diagram_dialect_flow.png")
    add_picture_safe(s11, diag_dialect, Inches(1.0), Inches(3.55), width=Inches(11.333))

    set_speaker_notes(s11, """ملاحظات المتحدث (Slide 11 - Dialect NLP in Action):
الرسم التوضيحي بالأسفل بيشرح خط سير معالجة الأوامر بالعامية المصرية:
التاجر بيكتب بلغته الطبيعية، ومحرك الـ NLP بيحول الكلمات العامية لـ Tokens مقننة ومطابقة لمخطط قواعد البيانات، ثم فحص هامش الربح، وأخيراً التنفيذ الآمن في قاعدة البيانات.""")

    # =========================================================================
    # SLIDE 12: SECTION 5 — ENTERPRISE TECH STACK & MONOREPO
    # =========================================================================
    s12 = prs.slides.add_slide(layout)
    set_slide_background(s12, BG_SLIDE)
    add_header(s12, "Section 5: Enterprise Tech Stack & Engineering Monorepo", "5. TECHNICAL ARCHITECTURE", "Production-grade multi-tenant stack engineered for high throughput, security, and low latency")
    add_footer(s12, 12)

    tiers = [
        ("1. Client & GenUI", PRIMARY_BLUE, PRIMARY_DARK, [
            ("React 18 & Vite", "Ultra-fast bundling and HMR."),
            ("Redux Toolkit", "Centralized merchant state."),
            ("Tailwind CSS", "Modern sleek styling system."),
            ("Recharts & Lucide", "Data visualization & SVG icons.")
        ]),
        ("2. Backend & Agent", AI_VIOLET, AI_VIOLET_DARK, [
            ("FastAPI & NestJS", "Async microservices & validation."),
            ("LangGraph Core", "Deterministic state machines."),
            ("sqlglot AST", "Abstract Syntax Tree security."),
            ("PromptGuard", "Adversarial injection defense.")
        ]),
        ("3. MLOps Engine", EMERALD_GREEN, EMERALD_DARK, [
            ("LightGBM Regressors", "Fast quantile gradient boosting."),
            ("SHAP TreeExplainer", "Feature impact attribution."),
            ("Scikit-Learn", "Log-RFM K-Means clustering."),
            ("MLflow & MinIO", "Model registry & artifact tracking.")
        ]),
        ("4. Data & Edge AI", AMBER_ORANGE, AMBER_DARK, [
            ("PostgreSQL 16 RLS", "Kernel session tenant isolation."),
            ("Redis 7 In-Memory", "Sub-millisecond query caching."),
            ("Meilisearch", "Typo-tolerant Arabic search."),
            ("Qwen2.5 4-Bit GGUF", "Air-gapped edge inference.")
        ])
    ]

    for i, (tier_title, color, dark_color, techs) in enumerate(tiers):
        x = Inches(0.8 + i * 2.98)
        add_modern_card(s12, x, Inches(1.55), Inches(2.8), Inches(5.2), BG_CARD, BORDER_LIGHT, top_stripe_color=color)

        pill = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.18), Inches(1.75), Inches(2.44), Inches(0.38))
        pill.fill.solid()
        pill.fill.fore_color.rgb = BG_CARD_SUBTLE
        pill.line.color.rgb = color
        pill.line.width = Pt(1)
        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = tier_title.upper()
        p_p.font.name = "Segoe UI"
        p_p.font.size = Pt(9)
        p_p.font.bold = True
        p_p.font.color.rgb = dark_color
        p_p.alignment = PP_ALIGN.CENTER

        tb = s12.shapes.add_textbox(x + Inches(0.18), Inches(2.3), Inches(2.44), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        for name, desc in techs:
            p_n = tf.add_paragraph()
            p_n.text = f"✔ {name}"
            p_n.font.name = "Segoe UI"
            p_n.font.size = Pt(10)
            p_n.font.bold = True
            p_n.font.color.rgb = TEXT_HEADING
            p_n.space_before = Pt(8)

            p_d = tf.add_paragraph()
            p_d.text = desc
            p_d.font.name = "Segoe UI"
            p_d.font.size = Pt(8.5)
            p_d.font.color.rgb = TEXT_BODY
            p_d.space_before = Pt(1)

    set_speaker_notes(s12, """ملاحظات المتحدث (Slide 12 - Tech Stack):
معمارية النظام مبنية بأحدث معايير هندسة البرمجيات وعلوم البيانات:
الواجهة الأمامية تعتمد على React 18 وVite، والباك إند مبني بـ FastAPI وLangGraph مع جدار حماية AST، وخادم تعلم الآلة مدعوم بـ LightGBM وSHAP، وقاعدة البيانات PostgreSQL 16 RLS مع كاش Redis فائق السرعة.""")

    # =========================================================================
    # SLIDE 13: SECTION 5 — OFFLINE & EDGE AI: QUANTIZED GGUF MODEL
    # =========================================================================
    s13 = prs.slides.add_slide(layout)
    set_slide_background(s13, BG_SLIDE)
    add_header(s13, "Offline & Edge AI: Quantized GGUF Model Runtime", "5. AIR-GAPPED & EDGE AI", "Zero-cloud, zero-API, completely private local execution on consumer hardware", badge_color=EMERALD_DARK, badge_bg=EMERALD_BG, badge_border=EMERALD_GREEN)
    add_footer(s13, 13)

    # Left: Edge Architecture Cards
    add_modern_card(s13, Inches(0.8), Inches(1.55), Inches(4.5), Inches(5.2), BG_CARD, BORDER_LIGHT, top_stripe_color=EMERALD_GREEN)
    tb_m = s13.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(4.1), Inches(4.8))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "Fine-Tuned Qwen2.5-Coder-3B"
    p.font.name = "Segoe UI"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD_DARK

    m_features = [
        ("QLoRA Dialect Fine-Tuning", "Base model fine-tuned on custom Egyptian retail instructions and SQL schemas."),
        ("4-Bit K-Quantization (Q4_K_M)", "Converted via llama.cpp, preserving 97.4% of full-precision accuracy."),
        ("Air-Gapped Privacy & Security", "Zero internet access needed; merchant financial data never leaves premises."),
        ("Zero API Recurring Cost", "Eliminates per-token cloud costs, ensuring viability for Egyptian SMEs.")
    ]
    for ft_title, ft_desc in m_features:
        p_ft = tf_m.add_paragraph()
        p_ft.text = f"✔ {ft_title}"
        p_ft.font.name = "Segoe UI"
        p_ft.font.size = Pt(10)
        p_ft.font.bold = True
        p_ft.font.color.rgb = TEXT_HEADING
        p_ft.space_before = Pt(8)

        p_fd = tf_m.add_paragraph()
        p_fd.text = ft_desc
        p_fd.font.name = "Segoe UI"
        p_fd.font.size = Pt(8.5)
        p_fd.font.color.rgb = TEXT_BODY
        p_fd.space_before = Pt(1)

    # Right: Hardware Benchmarks Picture
    add_modern_card(s13, Inches(5.5), Inches(1.55), Inches(7.033), Inches(5.2), BG_CARD, BORDER_LIGHT)
    chart_hw = os.path.join(ASSETS_DIR, "chart_hardware_bench.png")
    add_picture_safe(s13, chart_hw, Inches(5.6), Inches(1.75), width=Inches(6.833))

    set_speaker_notes(s13, """ملاحظات المتحدث (Slide 13 - Offline GGUF Model):
إنجاز تقني بارز في مشروع التخرج هو قدرتنا على تشغيل المنصة بالكامل Offline وبدون أي تكاليف سحابية:
الرسومات البيانية على اليمين توضح:
1- استهلاك الذاكرة العشوائية VRAM: فقط 2.18 جيجا على كارت شاشة اقتصادي RTX 3050 (أقل بكثير من سقف الـ 4 جيجا).
2- سرعة التوليد: 38.4 توكن في الثانية، مما يوفر استجابة فورية للتاجر مع سرية كاملة لبياناته المالية.""")

    # =========================================================================
    # SLIDE 14: SECTION 5 — ACADEMIC BENCHMARKS & QA EVALUATION
    # =========================================================================
    s14 = prs.slides.add_slide(layout)
    set_slide_background(s14, BG_SLIDE)
    add_header(s14, "Experimental Evaluation & Academic Benchmarks", "5. RIGOROUS EVALUATION & QA", "Empirical validation across a 100-prompt academic benchmark, unit suites, and latency tests", badge_color=EMERALD_DARK, badge_bg=EMERALD_BG, badge_border=EMERALD_GREEN)
    add_footer(s14, 14)

    # Top KPI Stat Boxes
    kpis = [
        ("98.0%", "Tool Selection F1-Score", "10 functional domain intent routes", PRIMARY_BLUE, PRIMARY_DARK),
        ("92.4%", "SQL Execution (EX)", "Identical result set to ground truth", AI_VIOLET, AI_VIOLET_DARK),
        ("99.2%", "Valid SQL Rate (VSR)", "Queries parsing without syntax break", EMERALD_GREEN, EMERALD_DARK),
        ("100.0%", "Security Block Rate", "Zero injection or DDL bypasses", ROSE_RED, ROSE_DARK)
    ]
    for i, (metric, label, sub, color, dark_color) in enumerate(kpis):
        x = Inches(0.8 + i * 2.98)
        add_modern_card(s14, x, Inches(1.55), Inches(2.8), Inches(1.35), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s14.shapes.add_textbox(x + Inches(0.15), Inches(1.65), Inches(2.5), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p_m = tf.paragraphs[0]
        p_m.text = metric
        p_m.font.name = "Segoe UI"
        p_m.font.size = Pt(24)
        p_m.font.bold = True
        p_m.font.color.rgb = dark_color
        p_m.alignment = PP_ALIGN.CENTER

        p_l = tf.add_paragraph()
        p_l.text = label
        p_l.font.name = "Segoe UI"
        p_l.font.size = Pt(9.5)
        p_l.font.bold = True
        p_l.font.color.rgb = TEXT_HEADING
        p_l.alignment = PP_ALIGN.CENTER

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(8)
        p_s.font.color.rgb = TEXT_MUTED
        p_s.alignment = PP_ALIGN.CENTER

    # Bottom Benchmarks Picture
    add_modern_card(s14, Inches(0.8), Inches(3.05), Inches(11.733), Inches(3.7), BG_CARD, BORDER_LIGHT)
    chart_bench = os.path.join(ASSETS_DIR, "chart_benchmarks.png")
    add_picture_safe(s14, chart_bench, Inches(1.0), Inches(3.15), width=Inches(11.333))

    set_speaker_notes(s14, """ملاحظات المتحدث (Slide 14 - Evaluation & Benchmarks):
لإثبات الجدارة العلمية للمشروع، نفذنا تقييماً أكاديمياً صارماً على بنشمارك مكون من 100 برومبت معقد:
الرسم على اليسار يوضح تفوق نموذج التنبؤ بنسبة خطأ WAPE 14.2% مقارنة بـ Prophet (28.6%) وARIMA (34.1%).
والرسم على اليمين يبرز دقة تصنيف النوايا 98%، ودقة تنفيذ الـ SQL 92.4%، ونسبة حماية 100% ضد هجمات الـ Injection.""")

    # =========================================================================
    # SLIDE 15: SECTION 6 — FUTURE ROADMAP & STRATEGIC VISION
    # =========================================================================
    s15 = prs.slides.add_slide(layout)
    set_slide_background(s15, BG_SLIDE)
    add_header(s15, "Section 6: Future Roadmap & Research Extensions", "6. FUTURE STRATEGY & ROADMAP", "Scaling from conversational commerce to autonomous social agents and e-invoicing compliance")
    add_footer(s15, 15)

    # Top Highlights
    future_highlights = [
        ("ETA E-Invoicing SDK", "Automated cryptographic signing with Egyptian Tax Authority.", PRIMARY_BLUE, PRIMARY_DARK, PRIMARY_BG),
        ("Egyptian Voice AI (Whisper)", "Fine-tuned ASR on regional Egyptian accents for hands-free commerce.", EMERALD_GREEN, EMERALD_DARK, EMERALD_BG),
        ("Autonomous Social Co-Pilot", "Live order capturing and bidding on Instagram & TikTok video streams.", AI_VIOLET, AI_VIOLET_DARK, AI_VIOLET_BG)
    ]
    for i, (title, sub, color, dark_color, bg_pill) in enumerate(future_highlights):
        x = Inches(0.8 + i * 3.98)
        add_modern_card(s15, x, Inches(1.55), Inches(3.78), Inches(0.9), BG_CARD, BORDER_LIGHT, top_stripe_color=color)
        tb = s15.shapes.add_textbox(x + Inches(0.15), Inches(1.65), Inches(3.48), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = dark_color

        p_s = tf.add_paragraph()
        p_s.text = sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(8.5)
        p_s.font.color.rgb = TEXT_MUTED

    # Bottom Roadmap Diagram Picture
    add_modern_card(s15, Inches(0.8), Inches(2.6), Inches(11.733), Inches(4.15), BG_CARD, BORDER_LIGHT)
    diag_road = os.path.join(ASSETS_DIR, "diagram_roadmap_timeline.png")
    add_picture_safe(s15, diag_road, Inches(1.0), Inches(2.7), width=Inches(11.333))

    set_speaker_notes(s15, """ملاحظات المتحدث (Slide 15 - Future Roadmap):
الرسم البياني بالأسفل يوضح خريطة طريق تطوير المنصة بعد التخرج عبر 4 مراحل استراتيجية:
المرحلة الأولى: الأساس المتين (اكتمل بالفعل).
المرحلة الثانية: التكامل مع منظومة الفاتورة الإلكترونية لمصلحة الضرائب المصرية ومساعد صوتي باللهجة العامية.
المرحلة الثالثة: وكلاء أذكياء لإدارة البث المباشر على تيك توك وإنستجرام.
المرحلة الرابعة: التوسع الإقليمي وأنظمة التعلم الموحد.""")

    # =========================================================================
    # SLIDE 16: CONCLUSION & TEAM ACKNOWLEDGMENT (Outro)
    # =========================================================================
    s16 = prs.slides.add_slide(layout)
    set_slide_background(s16, BG_SLIDE)
    add_header(s16, "Conclusion: The Next Frontier of Intelligent Commerce", "6. CONCLUSION & TEAM OUTRO", "Empowering Egyptian & MENA merchants through autonomous data science and AI agents")
    add_footer(s16, 16)

    # Left: Key Takeaways
    add_modern_card(s16, Inches(0.8), Inches(1.55), Inches(5.5), Inches(5.2), BG_CARD, BORDER_LIGHT, top_stripe_color=PRIMARY_BLUE)
    tb_c = s16.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.1), Inches(4.8))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "Key Project Achievements"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    takeaways = [
        ("Full-Stack Monorepo", "React 18 GenUI, FastAPI, LangGraph, and PostgreSQL 16 RLS."),
        ("Deterministic HITL Safety", "Pioneered dry-run preview cards eliminating financial risk."),
        ("Conversational BI", "AST-sandboxed Text-to-SQL engine across 71+ relational tables."),
        ("Predictive MLOps", "Quantile LightGBM models with 95% CI and SHAP explainability."),
        ("Edge Air-Gapped AI", "Runs locally on RTX 3050 (2.18 GB VRAM) with zero API costs.")
    ]
    for tk_title, tk_desc in takeaways:
        p_tk = tf_c.add_paragraph()
        p_tk.text = f"★ {tk_title}"
        p_tk.font.name = "Segoe UI"
        p_tk.font.size = Pt(10)
        p_tk.font.bold = True
        p_tk.font.color.rgb = TEXT_HEADING
        p_tk.space_before = Pt(8)

        p_td = tf_c.add_paragraph()
        p_td.text = tk_desc
        p_td.font.name = "Segoe UI"
        p_td.font.size = Pt(8.5)
        p_td.font.color.rgb = TEXT_BODY
        p_td.space_before = Pt(1)

    # Right: Team & Live Demo Callout
    add_modern_card(s16, Inches(6.5), Inches(1.55), Inches(6.033), Inches(5.2), BG_CARD, BORDER_LIGHT, top_stripe_color=EMERALD_GREEN)
    tb_ro = s16.shapes.add_textbox(Inches(6.75), Inches(1.7), Inches(5.5), Inches(4.8))
    tf_ro = tb_ro.text_frame
    tf_ro.word_wrap = True

    p2 = tf_ro.paragraphs[0]
    p2.text = "Data Science Graduation Team & Q&A"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = EMERALD_DARK

    p_names = tf_ro.add_paragraph()
    p_names.text = "فريق العمل (قسم علوم البيانات والذكاء الاصطناعي):"
    p_names.font.name = "Segoe UI"
    p_names.font.size = Pt(10.5)
    p_names.font.bold = True
    p_names.font.color.rgb = TEXT_HEADING
    p_names.space_before = Pt(6)

    team_members_ar = [
        ("مصطفى ماهر", "قائد الفريق، معمارية الوكلاء الذكية وعزل المستأجرين"),
        ("عادل أحمد", "تطوير الواجهات الأمامية وبطاقات Generative UI ومحرك الكتالوج"),
        ("عبد الله سامح", "محرك Text-to-SQL وجدار الحماية AST واستعلامات الـ ERP"),
        ("ريتاج حكيم", "هندسة الـ MLOps والتنبؤ بالطلب LightGBM وتفسيرات SHAP"),
        ("رناد حكيم", "تجزئة العملاء Log-RFM والتكامل مع بوابات Paymob وشحن بوسطة"),
        ("رُؤى أشرف", "إدارة الذاكرة ثلاثية المستويات ومركز القيادة وجدار الحماية")
    ]
    for m_ar, role_ar in team_members_ar:
        p_m = tf_ro.add_paragraph()
        p_m.text = f"• {m_ar} — {role_ar}"
        p_m.font.name = "Segoe UI"
        p_m.font.size = Pt(8.5)
        p_m.font.color.rgb = TEXT_BODY
        p_m.space_before = Pt(3)

    qa_pill = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(5.75), Inches(5.5), Inches(0.75))
    qa_pill.fill.solid()
    qa_pill.fill.fore_color.rgb = PRIMARY_BG
    qa_pill.line.color.rgb = PRIMARY_BLUE
    qa_pill.line.width = Pt(1.5)
    tf_qa = qa_pill.text_frame
    p_qa1 = tf_qa.paragraphs[0]
    p_qa1.text = "THANK YOU! QUESTIONS & LIVE DEMO"
    p_qa1.font.name = "Segoe UI"
    p_qa1.font.size = Pt(11)
    p_qa1.font.bold = True
    p_qa1.font.color.rgb = PRIMARY_DARK
    p_qa1.alignment = PP_ALIGN.CENTER

    p_qa2 = tf_qa.add_paragraph()
    p_qa2.text = "Storefront: http://localhost:5173  |  FastAPI Docs: http://localhost:8001/docs"
    p_qa2.font.name = "Segoe UI"
    p_qa2.font.size = Pt(8.5)
    p_qa2.font.color.rgb = TEXT_HEADING
    p_qa2.alignment = PP_ALIGN.CENTER
    p_qa2.space_before = Pt(1)

    set_speaker_notes(s16, """ملاحظات المتحدث (Slide 16 - Conclusion & Outro):
في الختام، مشروع توك شو (TalkShow) بيبرهن على إمكانية تطويع تقنيات الذكاء الاصطناعي التوليدي، ونماذج تعلم الآلة التنبؤية، لحل مشاكل تجارية حقيقية وملموسة في السوق المصري والمنطقة العربية بأعلى درجات الأمان والموثوقية.
بنشكر أساتذتنا المشرفين في قسم علوم البيانات، ولجنة التحكيم الموقرة على وقتهم وملاحظاتهم القيمة.
ودلوقتي بيسعدنا تلقي استفسارات وأسئلة حضراتكم وبندعوكم لبدء العرض الحي المباشر (Live Demo) على المنصة.""")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(OUTPUT_PRES_DIR, exist_ok=True)
    prs.save(OUTPUT_FILE)
    shutil.copy2(OUTPUT_FILE, OUTPUT_PRES_FILE)
    print(f"TalkShow presentation saved successfully to:\n1) {OUTPUT_FILE}\n2) {OUTPUT_PRES_FILE}")

if __name__ == "__main__":
    build_presentation()
