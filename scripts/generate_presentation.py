import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Color Palette & Theme Constants (Modern AI / Data Science Tech Theme)
# -----------------------------------------------------------------------------
BG_DARK = RGBColor(11, 15, 25)        # #0B0F19 Very deep slate navy
BG_CARD = RGBColor(22, 32, 50)        # #162032 Card container background
BG_CARD_LIGHT = RGBColor(30, 41, 59)  # #1E293B Slightly lighter card
BORDER_CYAN = RGBColor(14, 165, 233)  # #0EA5E9 Sky Cyan
BORDER_MUTED = RGBColor(51, 65, 85)   # #334155 Slate border

TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC Heading text
TEXT_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0 Body text
TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 Secondary / captions
TEXT_DIM = RGBColor(100, 116, 139)    # #64748B Subtext / footnotes

CYAN = RGBColor(14, 165, 233)         # #0EA5E9 TechZone Primary
CYAN_BRIGHT = RGBColor(56, 189, 248)  # #38BDF8 Bright Cyan
EMERALD = RGBColor(16, 185, 129)      # #10B981 Success / Metric / Green
PURPLE = RGBColor(139, 92, 246)       # #8B5CF6 AI / Agentic Violet
AMBER = RGBColor(245, 158, 11)        # #F59E0B Warning / Note
ROSE = RGBColor(244, 63, 94)          # #F43F5E Danger / Problem Accent

LOGO_PATH = r"D:\HITU\2\AI\Project-AI\logo\TransparentLogo.png"
OUTPUT_DIR = r"D:\HITU\2\AI\Project-AI\Docs"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "TechZone_Graduation_Presentation.pptx")
OUTPUT_PRES_FILE = os.path.join(r"D:\HITU\2\AI\Project-AI\Docs\Presentation", "TechZone_Graduation_Presentation.pptx")

# -----------------------------------------------------------------------------
# Presentation Initialization (16:9 Widescreen)
# -----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

def set_slide_background(slide, color=BG_DARK):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title, category_badge="TECHZONE ARCHITECTURE", subtitle=None):
    # Top logo
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, Inches(11.2), Inches(0.35), height=Inches(0.65))
        except Exception as e:
            pass
    
    # Category badge pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(3.2), Inches(0.32))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(20, 35, 60)
    pill.line.color.rgb = CYAN
    pill.line.width = Pt(1)
    tf_pill = pill.text_frame
    tf_pill.word_wrap = True
    tf_pill.margin_left = tf_pill.margin_top = tf_pill.margin_bottom = tf_pill.margin_right = 0
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = f"●  {category_badge.upper()}"
    p_pill.font.name = "Segoe UI"
    p_pill.font.size = Pt(9.5)
    p_pill.font.bold = True
    p_pill.font.color.rgb = CYAN_BRIGHT
    p_pill.alignment = PP_ALIGN.CENTER
    
    # Main Slide Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(10.2), Inches(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=BG_CARD, border_color=BORDER_MUTED, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_footer(slide, current_slide, total_slides=16):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = f"TechZone Platform | Graduation Project 2026 — Data Science Department                         Slide {current_slide} of {total_slides}"
    p.font.name = "Segoe UI"
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_DIM

def set_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

print("Presentation template framework ready.")
